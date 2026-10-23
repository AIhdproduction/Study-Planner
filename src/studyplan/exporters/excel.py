"""Excel workbook export for a computed StudyPlan result."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from studyplan.engine.rules import CheckResult
from studyplan.i18n import category_label, goal_label, localized_text, message, status_label
from studyplan.models import CreditsFile, Program

VALIDATION_LAST_ROW = 1000


def _header(sheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"


def _widths(sheet) -> None:
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            48, max(14, max(len(str(cell.value or "")) for cell in column) + 2)
        )


def _sheet_range(sheet_name: str, column: str, last_row: int) -> str:
    """Create an Excel-safe range formula, including localized sheet names."""
    escaped = sheet_name.replace("'", "''")
    return f"='{escaped}'!${column}$2:${column}${last_row}"


def export_workbook(
    path: Path, program: Program, credits: CreditsFile, result: CheckResult, language: str
) -> Path:
    """Write a reviewable workbook; calculated verdict values originate from the engine."""
    workbook = Workbook()
    overview = workbook.active
    overview.title = message(language, "overview")
    overview.append([message(language, "overall_verdict"), status_label(language, result.verdict)])
    overview.append([message(language, "computed"), date.today().isoformat()])
    overview.append([])
    overview.append([message(language, "goals")])
    for goal in result.goals:
        overview.append([goal_label(language, goal.id), status_label(language, goal.verdict)])
    overview.append([])
    overview.append([message(language, "category_ects")])
    targets = {
        rule.parameters.category_id: rule.parameters.ects
        for rule in program.rules
        if rule.parameters.type == "min_ects_in_category"
    }
    for category in program.categories:
        overview.append(
            [
                category_label(language, category.official_name),
                f"{result.category_ects.get(category.id, 0)} / {targets.get(category.id, 0)}",
            ]
        )
    overview.append(
        [message(language, "total_ects"), f"{result.total_ects} / {program.total_ects}"]
    )
    overview.append([])
    overview.append([message(language, "disclaimer")])
    overview.column_dimensions["A"].width = 80

    rules = workbook.create_sheet(message(language, "rules"))
    _header(
        rules,
        [
            message(language, "id"),
            message(language, "type"),
            message(language, "status"),
            message(language, "current"),
            message(language, "target"),
            message(language, "gap"),
            message(language, "provisional_reason"),
        ],
    )
    for rule in result.rules:
        rules.append(
            [
                rule.id,
                rule.type,
                status_label(language, rule.status),
                str(rule.current_value),
                str(rule.target_value),
                localized_text(language, rule.gap.text) if rule.gap else "",
                ", ".join(localized_text(language, reason) for reason in rule.provisional_reasons),
            ]
        )
    for status, color in {
        "fulfilled": "C6EFCE",
        "not_fulfilled": "FFC7CE",
        "provisional": "FCE4D6",
        "manual": "D9E1F2",
    }.items():
        rules.conditional_formatting.add(
            "C2:C1048576",
            CellIsRule(
                operator="equal",
                formula=[f'"{status_label(language, status)}"'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )

    modules = workbook.create_sheet(message(language, "modules"))
    _header(
        modules,
        [
            message(language, "code"),
            message(language, "title"),
            message(language, "ects"),
            message(language, "categories"),
            message(language, "mandatory"),
            "",
            message(language, "category"),
        ],
    )
    for index, item in enumerate(program.modules, start=2):
        modules.cell(index, 1, item.code)
        modules.cell(index, 2, item.title)
        modules.cell(index, 3, item.ects)
        modules.cell(index, 4, ", ".join(item.eligible_category_ids))
        modules.cell(index, 5, item.mandatory)
    for index, category in enumerate(program.categories, start=2):
        modules.cell(index, 7, category.id)

    sheet = workbook.create_sheet(message(language, "credits"))
    _header(
        sheet,
        [
            message(language, "module_code"),
            message(language, "semester"),
            message(language, "ects"),
            message(language, "grade"),
            message(language, "passed"),
            message(language, "category"),
            message(language, "recognition_status"),
        ],
    )
    allocations = {
        item.module_code: item.allocated_category_id
        for item in result.allocation
        if item.module_code
    }
    for item in credits.credits:
        sheet.append(
            [
                item.module_code,
                str(item.semester),
                item.ects,
                item.grade,
                item.passed,
                allocations.get(item.module_code),
                "",
            ]
        )
    for item in credits.recognized_credits:
        sheet.append(
            [
                item.module_code,
                "",
                item.ects,
                "",
                True,
                item.category_id,
                status_label(language, item.status),
            ]
        )
    table = Table(displayName="MyCredits", ref=f"A1:G{VALIDATION_LAST_ROW}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    code_validation = DataValidation(
        type="list",
        formula1=_sheet_range(modules.title, "A", len(program.modules) + 1),
        allow_blank=True,
    )
    category_validation = DataValidation(
        type="list",
        formula1=_sheet_range(modules.title, "G", len(program.categories) + 1),
        allow_blank=True,
    )
    sheet.add_data_validation(code_validation)
    sheet.add_data_validation(category_validation)
    code_validation.add(f"A2:A{VALIDATION_LAST_ROW}")
    category_validation.add(f"F2:F{VALIDATION_LAST_ROW}")
    summary_row = VALIDATION_LAST_ROW + 2
    sheet.cell(summary_row, 1, message(language, "category"))
    sheet.cell(summary_row, 2, message(language, "category_formula"))
    for index, category in enumerate(program.categories, start=summary_row + 1):
        sheet.cell(index, 1, category_label(language, category.official_name))
        sheet.cell(
            index,
            2,
            f'=SUMIF(F2:F{VALIDATION_LAST_ROW},"{category.id}",C2:C{VALIDATION_LAST_ROW})',
        )

    allocation = workbook.create_sheet(message(language, "allocation"))
    _header(
        allocation,
        [
            message(language, "module"),
            message(language, "recognized_credit"),
            message(language, "category"),
            message(language, "ects"),
            message(language, "reason"),
        ],
    )
    for item in result.allocation:
        allocation.append(
            [
                item.module_code,
                item.recognized_credit_id,
                item.allocated_category_id,
                item.ects_counted,
                localized_text(language, item.reason),
            ]
        )
    gaps = workbook.create_sheet(message(language, "gaps"))
    _header(gaps, [message(language, "id"), message(language, "status"), message(language, "text")])
    for gap in result.gaps:
        gaps.append(
            [
                gap["rule_id"],
                status_label(language, gap["status"]),
                localized_text(language, gap["text"]),
            ]
        )
    sources = workbook.create_sheet(message(language, "sources"))
    _header(
        sources,
        [message(language, "id"), message(language, "document"), message(language, "page")],
    )
    for rule in program.rules:
        if rule.source:
            sources.append([rule.id, rule.source.document, rule.source.page])
    issues = workbook.create_sheet(message(language, "issues"))
    _header(
        issues,
        [message(language, "id"), message(language, "severity"), message(language, "text")],
    )
    for issue in result.open_issues:
        issues.append([issue.id, issue.severity, localized_text(language, issue.text)])
    for worksheet in workbook.worksheets:
        if worksheet.max_row > 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        _widths(worksheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
