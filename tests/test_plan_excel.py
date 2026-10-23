"""Study-plan sheet presentation tests."""

from __future__ import annotations

from openpyxl import load_workbook

from studyplan.engine.plan import condition_text, generate_plan, plan_report
from studyplan.engine.rules import evaluate
from studyplan.exporters.excel import export_workbook
from studyplan.models import CreditsFile, PlanFile, PlanItem, RecognizedCredit, SourceReference

from .engine_support import checklist, credit, credits, program, student


def test_plan_excel_has_sheets_colors_conditions_and_german_labels(tmp_path) -> None:
    selected, selected_student = program(), student()
    plan, _ = generate_plan(selected, selected_student, credits(), PlanFile(), checklist(), "HS99")
    next(item for item in plan.planned_modules if item.module_code == "SYN.CORE1").status = "passed"
    plan.planned_modules.append(
        PlanItem(
            module_code="SYN.ADVANCED",
            target_semester="HS00",
            status="planned",
            origin="manual",
        )
    )
    result = evaluate(selected, selected_student, credits(), plan, checklist(), "forecast")
    path = tmp_path / "plan.xlsx"
    export_workbook(path, selected, credits(), result, "de", plan, selected_student, checklist())
    workbook = load_workbook(path)
    assert {"Studienplan", "Semesterübersicht"} <= set(workbook.sheetnames)
    sheet = workbook["Studienplan"]
    assert sheet.freeze_panes == "A2" and sheet["A2"].fill.fgColor.rgb
    assert all("Kern-Project" not in str(cell.value) for row in sheet.iter_rows() for cell in row)
    rows = {row[1].value: row for row in sheet.iter_rows(min_row=2) if row[1].value}
    assert rows["SYN.CORE1"][10].value in {None, ""}
    assert "SYN.CORE2" in rows["SYN.ADVANCED"][10].value


def test_plan_excel_uses_attempt_grades_and_recognized_ects_in_semester_totals(tmp_path) -> None:
    selected, selected_student = program(), student()
    store = CreditsFile(
        credits=[
            credit("SYN.CORE1", "HS99", passed=False, attempt=1),
            credit("SYN.CORE1", "FS00", passed=True, attempt=2),
        ],
        recognized_credits=[
            RecognizedCredit(
                id="recognized",
                category_id="related",
                ects=3,
                basis="synthetic recognition",
                decision_document=SourceReference(document="recognition.pdf"),
                status="provisional",
            )
        ],
    )
    plan = PlanFile(
        planned_modules=[
            PlanItem(
                module_code="SYN.CORE1",
                target_semester="HS99",
                status="failed",
                origin="credits",
                attempt=1,
            ),
            PlanItem(
                module_code="SYN.CORE1",
                target_semester="FS00",
                status="passed",
                origin="credits",
                attempt=2,
            ),
        ]
    )
    result = evaluate(selected, selected_student, store, plan, checklist(), "forecast")
    path = tmp_path / "attempts.xlsx"
    export_workbook(path, selected, store, result, "de", plan, selected_student, checklist())
    sheet = load_workbook(path, data_only=False)["Studienplan"]
    rows = {
        (
            row[0].value,
            row[1].value,
        ): row
        for row in sheet.iter_rows(min_row=2)
        if row[1].value
    }
    assert rows[("HS99", "SYN.CORE1")][8].value == "F"
    assert rows[("FS00", "SYN.CORE1")][8].value == "B"
    assert any(str(row[2].value).startswith("Anerkennung:") for row in sheet.iter_rows(min_row=2))
    report = plan_report(selected, selected_student, store, plan, checklist())
    totals = {}
    for row in sheet.iter_rows(min_row=2):
        if row[0].value and row[7].value != "nicht bestanden":
            totals[row[0].value] = totals.get(row[0].value, 0) + (row[5].value or 0)
    assert totals["HS99"] == report["semesters"]["HS99"]["planned_ects"]


def test_level_gate_text_skips_modules_passed_later() -> None:
    adv_item = PlanItem(module_code="SYN.ADVANCED", target_semester="HS01", status="planned")
    c_file = credits(
        credit("SYN.CORE1", "HS99", passed=False, attempt=1),
        credit("SYN.CORE1", "FS00", passed=True, attempt=2),
    )
    p_file = PlanFile(
        planned_modules=[
            PlanItem(module_code="SYN.CORE2", target_semester="FS00", status="planned"),
            adv_item,
        ]
    )
    text_de = condition_text("de", adv_item, program(), c_file, p_file)
    text_en = condition_text("en", adv_item, program(), c_file, p_file)
    assert "SYN.CORE1" not in text_de
    assert "SYN.CORE2" in text_de
    assert "erfüllt, sofern bestanden" in text_de
    assert "met if passed" in text_en


def test_level_gate_text_lists_each_module_once() -> None:
    adv_item = PlanItem(module_code="SYN.ADVANCED", target_semester="HS01", status="planned")
    c_file = credits(credit("SYN.CORE2", "HS99", passed=True))
    p_file = PlanFile(
        planned_modules=[
            PlanItem(module_code="SYN.CORE1", target_semester="HS99", status="running"),
            PlanItem(module_code="SYN.CORE1", target_semester="FS00", status="retake"),
            adv_item,
        ]
    )
    text_de = condition_text("de", adv_item, program(), c_file, p_file)
    assert text_de.count("SYN.CORE1") == 1
    assert "SYN.CORE1 (Wiederholung FS00)" in text_de
    assert "HS99" not in text_de


def test_level_gate_text_met_when_all_basics_passed() -> None:
    adv_item = PlanItem(module_code="SYN.ADVANCED", target_semester="HS01", status="planned")
    c_file = credits(
        credit("SYN.CORE1", "HS99", passed=True),
        credit("SYN.CORE2", "FS00", passed=True),
    )
    p_file = PlanFile(planned_modules=[adv_item])
    text_de = condition_text("de", adv_item, program(), c_file, p_file)
    text_en = condition_text("en", adv_item, program(), c_file, p_file)
    assert "Erfüllt: alle Basic-Pflichtmodule bestanden" in text_de
    assert "Met: all Basic mandatory modules passed" in text_en


def test_failed_row_points_to_retake() -> None:
    failed_item = PlanItem(module_code="SYN.CORE1", target_semester="HS99", status="failed")
    retake_item = PlanItem(module_code="SYN.CORE1", target_semester="FS00", status="retake")
    p_file = PlanFile(planned_modules=[failed_item, retake_item])
    c_file = credits(credit("SYN.CORE1", "HS99", passed=False))
    text_de = condition_text("de", failed_item, program(), c_file, p_file)
    text_en = condition_text("en", failed_item, program(), c_file, p_file)
    assert text_de == "Wiederholung geplant: FS00"
    assert text_en == "Retake planned: FS00"


def test_failed_row_reports_later_pass() -> None:
    failed_item = PlanItem(
        module_code="SYN.CORE1",
        target_semester="HS99",
        status="failed",
        attempt=1,
    )
    p_file = PlanFile(
        planned_modules=[
            failed_item,
            PlanItem(
                module_code="SYN.CORE1",
                target_semester="FS00",
                status="passed",
                attempt=2,
            ),
        ]
    )
    c_file = credits(
        credit("SYN.CORE1", "HS99", passed=False, attempt=1),
        credit("SYN.CORE1", "FS00", passed=True, attempt=2),
    )
    text_de = condition_text("de", failed_item, program(), c_file, p_file)
    text_en = condition_text("en", failed_item, program(), c_file, p_file)
    assert text_de == "Wiederholung bestanden: FS00 (Versuch 2)"
    assert text_en == "Retake passed: FS00 (attempt 2)"
