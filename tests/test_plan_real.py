"""Private-fixture plan-generation test without exposing personal data."""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from studyplan.cli import app
from studyplan.io.yaml_store import ProfileRepository, load_model
from studyplan.models import PlanFile, Semester


@pytest.mark.real
def test_real_mycampus_import_generates_a_retake_and_respects_the_level_gate(
    tmp_path: Path,
) -> None:
    fixture = Path("tests/fixtures/real/Meine Anmeldungen _ Hochschule Luzern.html")
    if not fixture.exists():
        pytest.skip("Real MyCampus fixture is unavailable.")
    runner = CliRunner()
    common = ["--study-dir", str(tmp_path), "--profiles-dir", str(Path("profiles").resolve())]
    initialized = runner.invoke(
        app,
        [
            *common,
            "init",
            "--university",
            "hslu_ta",
            "--program",
            "digital_construction_architecture_2025",
            "--intake-year",
            "2025",
            "--study-mode",
            "part_time",
            "--start-semester",
            "HS25",
        ],
    )
    assert initialized.exit_code == 0, initialized.output
    imported = runner.invoke(app, [*common, "credits", "import-mycampus", str(fixture), "--apply"])
    assert imported.exit_code == 0, imported.output
    recognized_minor = runner.invoke(
        app,
        [
            *common,
            "credits",
            "recognize",
            "add",
            "--ects",
            "15",
            "--basis",
            "synthetic provisional recognition",
            "--document",
            "recognition.pdf",
            "--category",
            "minor",
            "--status",
            "provisional",
        ],
    )
    assert recognized_minor.exit_code == 0, recognized_minor.output
    recognized_uncategorized = runner.invoke(
        app,
        [
            *common,
            "credits",
            "recognize",
            "add",
            "--ects",
            "3",
            "--basis",
            "synthetic provisional recognition",
            "--document",
            "recognition.pdf",
            "--no-category",
            "--status",
            "provisional",
        ],
    )
    assert recognized_uncategorized.exit_code == 0, recognized_uncategorized.output
    generated = runner.invoke(app, [*common, "plan", "generate", "--current-semester", "HS26"])
    assert generated.exit_code == 0, generated.output
    plan = load_model(tmp_path / "plan.yaml", PlanFile)
    report = runner.invoke(app, [*common, "plan", "--json"])
    assert report.exit_code == 0, report.output
    semesters = json.loads(report.output)["semesters"]
    running = [item for item in semesters["HS26"]["items"] if item["status"] == "running"]
    assert {item["code"] for item in running} == {
        "TA.BA_B_WERK",
        "TA.BA_DATA1_MM",
        "TA.BA_DC_TECH_MM",
        "TA.BA_DT_PROGR",
        "TA.BA_TUT",
    }
    assert sum(item["ects"] for item in running) == 18
    assert (
        next(item for item in semesters["HS25"]["items"] if item["code"] == "TA.BA_PYTHON_MM")[
            "grade"
        ]
        == "F"
    )
    assert (
        next(item for item in semesters["FS26"]["items"] if item["code"] == "TA.BA_WIRK")["grade"]
        == "F"
    )
    retake = next(
        item
        for item in plan.planned_modules
        if item.module_code == "TA.BA_WIRK" and item.status == "retake"
    )
    assert str(retake.target_semester) == "FS27"
    advanced = [
        item
        for item in plan.planned_modules
        if item.module_code == "TA.BA_DATA3_MM" and item.status == "planned"
    ]
    assert not advanced or advanced[0].target_semester.sort_key > retake.target_semester.sort_key
    assert (
        max(item.target_semester.sort_key for item in plan.planned_modules)
        <= Semester("FS31").sort_key
    )
    assert any(item.category_id == "related" for item in plan.placeholders)
    assert not any(
        str(item.target_semester) in {"HS25", "FS26", "HS26"} for item in plan.placeholders
    )
    python_hs25 = next(
        item for item in semesters["HS25"]["items"] if item["code"] == "TA.BA_PYTHON_MM"
    )
    assert "Wiederholung bestanden: FS26" in python_hs25["condition"]
    ress_condition = next(
        item["condition"]
        for semester in semesters.values()
        for item in semester["items"]
        if item["code"] == "TA.BA_RESS"
    )
    assert "TA.BA_PYTHON_MM" not in ress_condition
    assert ress_condition.count("TA.BA_WIRK") == 1
    assert (
        Semester(json.loads(report.output)["earliest_thesis_admission_semester"]).sort_key
        <= Semester("FS29").sort_key
    )
    assert semesters["HS25"]["cumulative_ects"] >= 18
    assert Semester(json.loads(report.output)["earliest_thesis_admission_semester"]) == Semester(
        "FS28"
    )
    exported = runner.invoke(app, [*common, "--lang", "de", "excel", "--mode", "forecast"])
    assert exported.exit_code == 0, exported.output
    workbook = load_workbook(next((tmp_path / "output").glob("*.xlsx")), data_only=False)
    study_plan = workbook["Studienplan"]
    values = list(study_plan.iter_rows(min_row=2, values_only=True))
    python_failure = next(row for row in values if row[0] == "HS25" and row[1] == "TA.BA_PYTHON_MM")
    wirkung_failure = next(row for row in values if row[0] == "FS26" and row[1] == "TA.BA_WIRK")
    assert python_failure[8] == "F" and wirkung_failure[8] == "F"
    assert sum(1 for row in values if str(row[2]).startswith("Anerkennung:")) == 2
    calculated = {}
    for row in values:
        if row[0] and row[7] != "nicht bestanden":
            calculated[row[0]] = calculated.get(row[0], 0) + (row[5] or 0)
    assert sum(calculated.values()) == 180
    profile = ProfileRepository(Path("profiles")).program(
        "hslu_ta", "digital_construction_architecture_2025"
    )
    ects = {module.code: module.ects for module in profile.modules}
    entries = profile.model_timetables["part_time"]["HS"]
    limit = max(
        sum(
            ects.get(entry.module_code, 0)
            for entry in entries
            if entry.relative_semester == semester
        )
        for semester in range(1, 9)
    )
    assert (
        max(
            ects
            for semester, ects in calculated.items()
            if Semester(semester).sort_key >= Semester("HS26").sort_key
        )
        <= limit
    )
    overview = workbook["Semesterübersicht"]
    assert any(
        "Thesis frühestens: FS28" in str(row[5].value) for row in overview.iter_rows(min_row=2)
    )
    for sheet in (study_plan, overview):
        text = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
        assert not any(
            word in text
            for word in ["passed", "failed", "running", "planned", "retake", "placeholder"]
        )
