"""Workbook structure and localization tests."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

from studyplan.cli import DISCLAIMER, app


def _create_workbook(tmp_path: Path, language: str):
    runner = CliRunner()
    profiles = Path("tests/fixtures/synthetic/profiles").resolve()
    common = ["--study-dir", str(tmp_path), "--profiles-dir", str(profiles)]
    assert (
        runner.invoke(
            app,
            [
                *common,
                "init",
                "--university",
                "synthetic_u",
                "--program",
                "synthetic_program_2099",
                "--intake-year",
                "2099",
                "--study-mode",
                "part_time",
                "--start-semester",
                "HS99",
                "--lang",
                language,
            ],
        ).exit_code
        == 0
    )
    assert (
        runner.invoke(
            app,
            [
                *common,
                "credits",
                "add",
                "--code",
                "SYN.CORE1",
                "--semester",
                "HS99",
                "--ects",
                "3",
                "--grade",
                "B",
            ],
        ).exit_code
        == 0
    )
    result = runner.invoke(app, [*common, "excel"])
    assert result.exit_code == 0, result.output
    return load_workbook(next((tmp_path / "output").glob("*.xlsx")), data_only=False)


def test_excel_english_structure_and_full_validation_ranges(tmp_path: Path) -> None:
    workbook = _create_workbook(tmp_path, "en")
    assert workbook.sheetnames == [
        "Overview",
        "Rules",
        "Modules",
        "My Credits",
        "Allocation",
        "Gaps",
        "Sources",
        "Open Issues",
    ]
    sheet = workbook["My Credits"]
    assert sheet.tables["MyCredits"].ref == "A1:G1000" and sheet.freeze_panes == "A2"
    validations = list(sheet.data_validations.dataValidation)
    assert len(validations) == 2
    assert all(validation.formula1.startswith("=") for validation in validations)
    assert all(len(validation.formula1) <= 255 for validation in validations)
    assert {str(validation.sqref) for validation in validations} == {"A2:A1000", "F2:F1000"}
    formulas = [
        cell.value for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str)
    ]
    assert '=SUMIF(F2:F1000,"core",C2:C1000)' in formulas
    assert len(workbook["Rules"].conditional_formatting) > 0
    overview = " ".join(
        str(cell.value) for row in workbook["Overview"].iter_rows() for cell in row if cell.value
    )
    assert "Thesis admission" in overview and "Graduation" in overview and DISCLAIMER in overview


def test_excel_german_localizes_sheets_headers_and_verdicts(tmp_path: Path) -> None:
    workbook = _create_workbook(tmp_path, "de")
    assert workbook.sheetnames == [
        "Übersicht",
        "Regeln",
        "Module",
        "Meine Leistungen",
        "Zuteilung",
        "Lücken",
        "Quellen",
        "Offene Punkte",
    ]
    overview = " ".join(
        str(cell.value) for row in workbook["Übersicht"].iter_rows() for cell in row if cell.value
    )
    assert "Gesamtergebnis" in overview and "Bachelorarbeitszulassung" in overview
    assert "nicht zugelassen" in overview and "Nicht rechtsverbindlich" in overview
    rules = workbook["Regeln"]
    assert rules.cell(1, 3).value == "Status"
    assert rules.cell(1, 7).value == "Grund für Vorläufigkeit"
    assert any(cell.value == "nicht erfüllt" for cell in rules["C"])
