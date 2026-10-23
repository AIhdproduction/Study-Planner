"""LibreOffice parity test for the simple Excel category formulas."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from studyplan.engine.rules import evaluate
from studyplan.exporters.excel import export_workbook
from tests.engine_support import checklist, credit, credits, plan, program, student


def test_excel_category_formula_parity_with_engine(tmp_path: Path) -> None:
    """LibreOffice calculates exported SUMIF values independently from the engine."""
    office = shutil.which("soffice")
    if office is None:
        pytest.skip("LibreOffice is not installed")
    selected_program = program()
    selected_credits = credits(credit("SYN.CORE1"), credit("SYN.CORE2"))
    result = evaluate(selected_program, student(), selected_credits, plan(), checklist(), "current")
    source = tmp_path / "source.xlsx"
    export_workbook(source, selected_program, selected_credits, result, "en")
    output_dir = tmp_path / "calculated"
    output_dir.mkdir()
    completed = subprocess.run(
        [office, "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)],
        check=False,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert completed.returncode == 0, completed.stderr
    calculated = load_workbook(output_dir / source.name, data_only=True)["My Credits"]
    values = {
        calculated.cell(row=row, column=1).value: calculated.cell(row=row, column=2).value
        for row in range(1, calculated.max_row + 1)
    }
    assert values["Core"] == result.category_ects["core"]
